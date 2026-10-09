#!/usr/bin/env python3
"""In tóm tắt các đoạn liên tục Step/Cycle để kiểm tra schema CALCE."""

from __future__ import annotations

import argparse
import io
import zipfile
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def summarize(archive_path: Path, member_filter: str) -> None:
    with zipfile.ZipFile(archive_path) as archive:
        members = [name for name in archive.namelist() if member_filter in name]
        for member in members:
            workbook = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=True)
            try:
                sheet = next(
                    workbook[name]
                    for name in workbook.sheetnames
                    if name.startswith("Channel_") and name != "Channel_Chart"
                )
                rows = list(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
            header = rows[0]
            columns = {name: header.index(name) for name in (
                "Test_Time(s)", "Step_Index", "Cycle_Index", "Current(A)",
                "Voltage(V)", "Charge_Capacity(Ah)", "Discharge_Capacity(Ah)",
            )}
            values = []
            for row in rows[1:]:
                try:
                    values.append(tuple(float(row[columns[name]]) for name in columns))
                except (TypeError, ValueError):
                    continue
            array = np.asarray(values, dtype=np.float64)
            step = array[:, 1].astype(int)
            cycle = array[:, 2].astype(int)
            boundaries = np.flatnonzero(
                np.concatenate(([True], (step[1:] != step[:-1]) | (cycle[1:] != cycle[:-1])))
            )
            stops = np.concatenate((boundaries[1:], [array.shape[0]]))
            print(f"\n{archive_path.name}/{member}: rows={array.shape[0]}")
            print("run cycle step rows t_start t_stop I_min I_max V_min V_max Qc_end Qd_end")
            for run, (start, stop) in enumerate(zip(boundaries, stops)):
                part = array[start:stop]
                print(
                    f"{run:3d} {cycle[start]:5d} {step[start]:4d} {stop-start:5d} "
                    f"{part[0,0]:8.1f} {part[-1,0]:8.1f} "
                    f"{part[:,3].min():7.3f} {part[:,3].max():7.3f} "
                    f"{part[:,4].min():6.3f} {part[:,4].max():6.3f} "
                    f"{part[-1,5]:7.3f} {part[-1,6]:7.3f}"
                )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("archive", type=Path)
    parser.add_argument("--member", default="50SOC")
    args = parser.parse_args()
    summarize(args.archive, args.member)


if __name__ == "__main__":
    main()

